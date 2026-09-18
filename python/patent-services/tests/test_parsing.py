"""parse_docx round-trips: build docx fixtures with python-docx, parse them back."""

from __future__ import annotations

from docx import Document

from patent_services.parsing import parse_docx, parse_docx_to_file


def build_docx(path, paragraphs):
    document = Document()
    for kind, text in paragraphs:
        if kind == "h":
            document.add_heading(text, level=1)
        elif kind == "p":
            document.add_paragraph(text)
        else:
            document.add_paragraph("")
    document.save(path)


def test_explicit_heading_style_becomes_markdown_heading(tmp_path):
    docx_path = tmp_path / "template.docx"
    build_docx(docx_path, [("h", "背景技术"), ("p", "现有方案存在不足。"), ("p", "")])
    markdown = parse_docx(str(docx_path))
    assert "# 背景技术" in markdown
    assert "现有方案存在不足。" in markdown


def test_three_level_numbering_maps_to_heading_levels(tmp_path):
    docx_path = tmp_path / "numbered.docx"
    build_docx(docx_path, [
        ("p", "1. 技术领域"),
        ("p", "1.1 具体领域"),
        ("p", "1.1.1 细分说明"),
        ("p", "一、发明内容"),
        ("p", "（一）技术方案"),
        ("p", "第二章 附图说明"),
        ("p", "普通段落不升级"),
    ])
    markdown = parse_docx(str(docx_path))
    assert "# 1. 技术领域" in markdown
    assert "## 1.1 具体领域" in markdown
    assert "### 1.1.1 细分说明" in markdown
    assert "# 一、发明内容" in markdown
    assert "## （一）技术方案" in markdown
    assert "# 第二章 附图说明" in markdown
    assert "普通段落不升级" in markdown
    assert "\n# 普通段落" not in markdown


def test_blank_paragraphs_are_preserved_as_blank_lines(tmp_path):
    docx_path = tmp_path / "blank.docx"
    build_docx(docx_path, [("p", "第一段"), ("p", ""), ("p", "第二段")])
    markdown = parse_docx(str(docx_path))
    assert "第一段\n\n第二段" in markdown


def test_parse_docx_to_file_writes_markdown(tmp_path):
    docx_path = tmp_path / "in" / "template.docx"
    docx_path.parent.mkdir(parents=True)
    build_docx(docx_path, [("h", "名称"), ("p", "一种存储装置")])
    output = tmp_path / "out" / "reference" / "template.md"
    written = parse_docx_to_file(str(docx_path), str(output))
    assert written == str(output)
    assert "# 名称" in output.read_text(encoding="utf-8")
